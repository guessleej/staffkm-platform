"""文件處理核心 — 載入、分段、向量化"""
import io
from pathlib import Path
from typing import BinaryIO

import structlog

log = structlog.get_logger()


class DocumentProcessor:
    """支援 PDF、Word、Excel、Markdown、TXT 的文件載入器。"""

    LOADERS: dict[str, str] = {
        ".pdf": "_load_pdf",
        ".docx": "_load_docx",
        ".doc":  "_load_doc_legacy",   # OLE2 → antiword
        ".xlsx": "_load_excel",
        ".xls":  "_load_xls_legacy",   # v5.9.19: OLE2 舊版 Excel → xlrd
        # v5.11.x: ODF（政府/教育公文常用 OpenDocument）
        ".odt":  "_load_odt",          # ODF 文字文件
        ".ods":  "_load_ods",          # ODF 試算表
        ".csv":  "_load_csv",
        ".md":   "_load_text",
        ".txt":  "_load_text",
        ".html": "_load_text",
        # v5.9.22: 圖片 OCR
        ".png":  "_load_image",
        ".jpg":  "_load_image",
        ".jpeg": "_load_image",
        ".webp": "_load_image",
        ".tiff": "_load_image",
        ".bmp":  "_load_image",
        # v5.13: 音檔 → 地端 ASR 逐字稿（帶時間碼）
        ".mp3":  "_load_audio",
        ".wav":  "_load_audio",
        ".m4a":  "_load_audio",
        ".flac": "_load_audio",
        ".ogg":  "_load_audio",
        ".aac":  "_load_audio",
        # v5.13: 字幕檔 → 直接吃文字（零 ASR 成本）
        ".srt":  "_load_srt",
        ".vtt":  "_load_vtt",
    }

    # v5.9.22: OCR 語言 (繁中 + 簡中 + 英文)
    OCR_LANG = "chi_tra+chi_sim+eng"
    # PDF 純文字層抽出少於此字數 → 視為掃描件，觸發 OCR fallback
    OCR_PDF_TEXT_THRESHOLD = 20

    def __init__(
        self,
        *,
        ocr_engine: str | None = None,
        vision_model: str | None = None,
        vision_base_url: str | None = None,
        vision_api_key: str | None = None,
    ):
        """v5.11.x: OCR / vision 設定可由呼叫端覆寫（接 system_settings.default.vision）。

        全部為 None 時 fallback 用 app.config.settings 的環境變數值，維持向後相容。
        """
        self._ocr_engine_override = ocr_engine
        self._vision_model_override = vision_model
        self._vision_base_url_override = vision_base_url
        self._vision_api_key_override = vision_api_key

    @staticmethod
    def _decode_bytes(data: bytes) -> str:
        """文字位元組 → str：先試 UTF-8，失敗再偵測編碼。

        台灣公文 / 舊版 Excel 匯出常見 Big5 / CP950（甚至簡中 GB18030），
        直接用 UTF-8 解會炸成 mojibake，故偵測時優先試這些編碼。
        """
        if not data:
            return ""
        # 1. UTF-8（含 BOM）strict — 最常見、誤判率最低
        try:
            return data.decode("utf-8-sig")
        except UnicodeDecodeError:
            pass
        # 2. charset-normalizer 偵測 — 對 Big5 vs GB18030 的區辨力佳，
        #    避免「GB bytes 剛好是合法 Big5 → 解成亂碼」這類 false positive。
        detected_text: str | None = None
        detected_enc = ""
        try:
            from charset_normalizer import from_bytes
            best = from_bytes(data).best()
            if best is not None:
                detected_text = str(best)
                detected_enc = (best.encoding or "").lower().replace("-", "_")
        except ImportError:
            pass
        # 偵測為 CJK 編碼時可信（big5/cp950/gb* 互斥度高）→ 直接採用
        _CJK = {"big5", "big5hkscs", "cp950", "gb18030", "gbk", "gb2312", "hz", "euc_cn"}
        if detected_enc in _CJK and detected_text is not None:
            return detected_text
        # 3. 偵測沒把握或猜成 Latin 系（短 Big5 文本常見誤判）→
        #    明確優先試台灣公文 / 舊 Excel 常見的 Big5 / CP950
        for enc in ("big5", "cp950"):
            try:
                return data.decode(enc)
            except (UnicodeDecodeError, LookupError):
                continue
        # 4. 用偵測結果，否則 GB18030 保底 → 最終 utf-8 replace
        if detected_text is not None:
            return detected_text
        try:
            return data.decode("gb18030")
        except (UnicodeDecodeError, LookupError):
            return data.decode("utf-8", errors="replace")

    def load(self, file: BinaryIO, filename: str) -> str:
        ext = Path(filename).suffix.lower()
        # 以 magic bytes 偵測真實格式（副檔名常常與內容不符）
        head = file.read(8)
        file.seek(0)
        actual = self._sniff(head, ext)

        method_name = self.LOADERS.get(actual)
        if not method_name:
            raise ValueError(
                f"不支援的檔案格式：{ext}（內容偵測為 {actual or '未知'}）。"
                f"請另存為 .docx / .pdf / .txt 後再上傳。"
            )
        return getattr(self, method_name)(file)

    @staticmethod
    def _sniff(head: bytes, ext: str) -> str:
        """以 magic bytes 偵測實際檔案格式，副檔名作為輔助。"""
        # PDF: %PDF
        if head.startswith(b"%PDF"):
            return ".pdf"
        # v5.9.22: 圖片 magic bytes
        if head.startswith(b"\x89PNG"):
            return ".png"
        if head.startswith(b"\xff\xd8\xff"):           # JPEG
            return ".jpg"
        if head.startswith(b"RIFF") and ext == ".webp":
            return ".webp"
        if head.startswith(b"II*\x00") or head.startswith(b"MM\x00*"):  # TIFF
            return ".tiff"
        if head.startswith(b"BM"):                     # BMP
            return ".bmp"
        # v5.13: 音檔 magic bytes
        if head[:3] == b"ID3" or head[:2] in (b"\xff\xfb", b"\xff\xf3", b"\xff\xf2"):  # MP3
            return ".mp3"
        if head[:4] == b"RIFF" and ext == ".wav":      # WAV（RIFF…WAVE；WAVE 在 offset 8、以副檔名區分）
            return ".wav"
        if head[4:8] == b"ftyp":                        # MP4/M4A/MOV 容器
            return ".m4a"
        if head[:4] == b"fLaC":                         # FLAC
            return ".flac"
        if head[:4] == b"OggS":                         # OGG
            return ".ogg"
        # v5.13: 字幕 — VTT 有 WEBVTT 開頭；SRT 無 magic，靠副檔名
        if head.startswith(b"WEBVTT") or (head.startswith(b"\xef\xbb\xbfWEBV")):
            return ".vtt"
        if ext in (".srt", ".vtt", ".aac"):
            return ext
        # DOCX/XLSX 等 Office Open XML：PK ZIP 容器
        if head.startswith(b"PK\x03\x04") or head.startswith(b"PK\x05\x06"):
            # 同樣是 PK ZIP；ODF 與 OOXML 都在此，需以副檔名區分
            if ext in (".odt", ".ods", ".odp"):
                return ext
            if ext in (".xlsx",):
                return ".xlsx"
            return ".docx"
        # 舊版 OLE2 (.doc / .xls)：D0 CF 11 E0 A1 B1 1A E1
        if head.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
            # v5.9.19: 加 xlrd 支援，.xls 不再 raise
            if ext == ".xls":
                return ".xls"
            return ".doc"
        # 純文字 / Markdown / CSV 沒有固定 magic bytes
        if ext in (".txt", ".md", ".csv"):
            return ext
        # 已知文字內容 → 視為 txt
        try:
            head.decode("utf-8")
            return ".txt"
        except UnicodeDecodeError:
            return ext  # 最終回退到副檔名

    def _load_pdf(self, file: BinaryIO) -> str:
        """v5.9.22 混合: 先抽文字層；幾乎抽不到字 (掃描件) → OCR fallback。"""
        from pypdf import PdfReader
        data = file.read()
        file.seek(0)
        import io as _io
        text = ""
        try:
            reader = PdfReader(_io.BytesIO(data))
            text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception:
            text = ""
        # 文字層夠多 → 直接用（文字型 PDF，快且免費）
        if len(text.strip()) >= self.OCR_PDF_TEXT_THRESHOLD:
            return text
        # 掃描件 / 圖片型 PDF → OCR
        ocr_text = self._ocr_pdf(data)
        # 兩者取較長的（避免 OCR 失敗時丟掉原本少量文字）
        return ocr_text if len(ocr_text.strip()) > len(text.strip()) else text

    def _ocr_pdf(self, data: bytes) -> str:
        """v5.9.22/v5.9.23: 掃描 PDF → pdf2image 轉頁圖 → OCR（引擎依設定）。"""
        try:
            from pdf2image import convert_from_bytes
            from PIL import Image  # noqa
            import io as _io
        except ImportError:
            return ""
        try:
            # v5.12: 200→300 DPI（OCR 標準）；掃描 PDF 轉圖解析度是 OCR 準度最大單一變因
            images = convert_from_bytes(data, dpi=300)
        except Exception:
            return ""
        pages = []
        for img in images:
            buf = _io.BytesIO()
            img.save(buf, format="PNG")
            txt = self._ocr_image_bytes(buf.getvalue())
            if txt.strip():
                pages.append(txt)
        return "\n\n".join(pages)

    def _load_image(self, file: BinaryIO) -> str:
        """v5.13 Phase 1: 圖片 = OCR 抽字（tesseract/vision）＋ 看圖說話（vision LLM 語意描述），
        兩段合併。視覺描述失敗 → 自動退回只用 OCR（不會壞）。純圖像/圖表也能進 KB 被檢索。
        """
        from app.config import settings
        data = file.read()

        # ① OCR 抽字（依 OCR_ENGINE：tesseract 預設 / vision；本身已含 fallback）
        ocr_text = ""
        try:
            ocr_text = self._ocr_image_bytes(data).strip()
        except Exception as e:  # noqa: BLE001
            log.warning("image_ocr_failed", error=str(e)[:200])

        # ② 看圖說話（同一個 vision LLM、描述 prompt）— 每張都跑；失敗不阻斷
        desc = ""
        if settings.IMAGE_DESCRIBE_ENABLED:
            try:
                desc = self._vision_describe(self._preprocess_for_ocr(data, "vision")).strip()
            except Exception as e:  # noqa: BLE001
                log.warning("image_describe_failed", error=str(e)[:200])

        parts: list[str] = []
        if ocr_text:
            parts.append("【圖片文字】\n" + ocr_text)
        if desc:
            parts.append("【圖片內容描述】\n" + desc)
        if not parts:
            raise ValueError(
                "圖片 OCR 未辨識到文字，視覺描述也無法取得（純圖像 / 太模糊，且 vision 模型不可用）。"
            )
        return "\n\n".join(parts)

    # ── OCR 前處理（引擎感知）────────────────────────────────────────
    def _preprocess_for_ocr(self, img_bytes: bytes, engine: str) -> bytes:
        """v5.12: OCR 前處理提升準度。失敗回原圖（不阻斷）。
        - 共同：解析度太低 → LANCZOS 升採樣（OCR 對小字最敏感）。
        - tesseract（古典 OCR）：轉灰階 + autocontrast（受惠；不做激進二值化以免過頭）。
        - vision/glm-ocr（vision-LLM）：保留彩色自然影像（二值化反而傷它），只做輕度對比。
        """
        try:
            from PIL import Image, ImageOps
            import io as _io
            img = Image.open(_io.BytesIO(img_bytes))
            if img.mode not in ("RGB", "L"):
                img = img.convert("RGB")
            if img.width < 1600:   # 升採樣（300DPI A4 約 2480px，掃描件常更小）
                scale = 1600 / img.width
                img = img.resize((int(img.width * scale), int(img.height * scale)), Image.LANCZOS)
            if engine == "tesseract":
                img = ImageOps.autocontrast(ImageOps.grayscale(img), cutoff=1)
            else:                  # vision：保留彩色，僅輕度對比
                img = ImageOps.autocontrast(img, cutoff=1)
            out = _io.BytesIO()
            img.save(out, format="PNG")
            return out.getvalue()
        except Exception as e:  # noqa: BLE001
            log.warning("ocr_preprocess_failed", error=str(e)[:120])
            return img_bytes

    # ── OCR 引擎分派 ────────────────────────────────────────────────
    def _ocr_image_bytes(self, img_bytes: bytes) -> str:
        """依 settings.OCR_ENGINE 選引擎；vision 失敗可 fallback tesseract。"""
        from app.config import settings
        engine = (self._ocr_engine_override or settings.OCR_ENGINE or "tesseract").lower()
        img_bytes = self._preprocess_for_ocr(img_bytes, engine)   # v5.12: 前處理
        if engine == "vision":
            try:
                txt = self._ocr_vision(img_bytes)
                if txt.strip():
                    return txt
                log.warning("vision_ocr_empty_fallback_tesseract")
            except Exception as e:
                log.warning("vision_ocr_failed", error=str(e)[:200])
                if not settings.VISION_OCR_FALLBACK_TESSERACT:
                    raise ValueError(f"Vision OCR 失敗：{e}")
            # fallback
            return self._ocr_tesseract(img_bytes)
        return self._ocr_tesseract(img_bytes)

    def _ocr_tesseract(self, img_bytes: bytes) -> str:
        try:
            import pytesseract
            from PIL import Image
            import io as _io
        except ImportError:
            return ""
        try:
            img = Image.open(_io.BytesIO(img_bytes))
            return pytesseract.image_to_string(img, lang=self.OCR_LANG)
        except Exception as e:
            log.warning("tesseract_ocr_failed", error=str(e)[:200])
            return ""

    def _vision_chat(self, img_bytes: bytes, prompt: str, *, temperature: float = 0.0,
                     max_tokens: int = 2048) -> str:
        """v5.13: 共用的 vision LLM 呼叫。OCR 與「看圖說話」都走這條，只是 prompt 不同。

        - 地端 ollama（預設）→ **原生 /api/chat + think:False**：thinking 型模型（如 gemma4:e4b）
          經 OpenAI-compat 無法關思考、會把 token 花光回空 → 必須走原生 API 帶 think:False。
        - 雲端 vision（VISION_USE_OLLAMA_NATIVE=false）→ OpenAI-compat /chat/completions。
        - **max_tokens 必設**：不設時一直生成，慢模型(glm-ocr bf16)破 timeout。
        """
        import base64, httpx
        from app.config import settings

        b64 = base64.b64encode(img_bytes).decode("ascii")
        # 覆寫優先（system_settings.default.vision）→ 否則 env settings
        vision_model = self._vision_model_override or settings.VISION_OCR_MODEL
        vision_base = self._vision_base_url_override or settings.VISION_OCR_BASE_URL or ""
        vision_key = self._vision_api_key_override or settings.VISION_OCR_API_KEY
        headers = {"Content-Type": "application/json"}
        if vision_key:
            headers["Authorization"] = f"Bearer {vision_key}"

        if settings.VISION_USE_OLLAMA_NATIVE:
            # ollama 原生：base 去掉尾端 /v1 → /api/chat；images 傳純 base64（無 data: 前綴）
            base = vision_base.rstrip("/")
            if base.endswith("/v1"):
                base = base[:-3].rstrip("/")
            url = f"{base}/api/chat"
            payload = {
                "model": vision_model,
                "messages": [{"role": "user", "content": prompt, "images": [b64]}],
                "stream": False,
                "think": False,   # 關 thinking：否則 gemma4 等思考型把 token 花光回空
                "options": {"num_predict": max_tokens, "temperature": temperature},
            }
            with httpx.Client(timeout=120.0) as client:
                resp = client.post(url, json=payload, headers=headers)
            if resp.status_code != 200:
                raise RuntimeError(f"HTTP {resp.status_code}: {resp.text[:200]}")
            return (resp.json().get("message") or {}).get("content", "") or ""

        # 雲端 OpenAI-compat
        url = f"{vision_base.rstrip('/')}/chat/completions"
        payload = {
            "model": vision_model,
            "messages": [{
                "role": "user",
                "content": [
                    {"type": "text", "text": prompt},
                    {"type": "image_url", "image_url": {"url": f"data:image/png;base64,{b64}"}},
                ],
            }],
            "temperature": temperature,
            "max_tokens": max_tokens,
            "stream": False,
        }
        with httpx.Client(timeout=120.0) as client:
            resp = client.post(url, json=payload, headers=headers)
        if resp.status_code != 200:
            raise RuntimeError(f"HTTP {resp.status_code}: {resp.text[:200]}")
        data = resp.json()
        return (data.get("choices") or [{}])[0].get("message", {}).get("content", "") or ""

    def _ocr_vision(self, img_bytes: bytes) -> str:
        """v5.9.23: Vision LLM OCR — 逐字提取圖片文字。"""
        return self._vision_chat(
            img_bytes,
            "你是 OCR 引擎。請逐字提取這張圖片中的所有文字，"
            "原樣輸出（保留換行 / 表格結構 / 完整保留所有標點符號），"
            "不要漏字、不要翻譯、不要加任何說明、標題或評論。"
            "若圖片沒有文字就回空字串。",
            temperature=0.0,
            max_tokens=4096,        # 整頁 OCR 文字可能很長
        )

    def _vision_describe(self, img_bytes: bytes) -> str:
        """v5.13 Phase 1: 看圖說話 — vision LLM 語意描述圖片內容（非逐字 OCR）。"""
        from app.config import settings
        return self._vision_chat(
            img_bytes,
            "用繁體中文簡要描述這張圖片，讓沒看到圖的人也能理解。"
            "說明：主題、人物/動作、場景、重要物件；若有圖表/流程圖則說明其結構與意義。"
            "條列 3~6 點、客觀、不臆測、不重複圖中已可逐字辨識的純文字。",
            temperature=0.2,
            max_tokens=settings.IMAGE_DESCRIBE_MAX_TOKENS,   # 上限避免慢模型一直生成破 timeout
        )

    def _load_docx(self, file: BinaryIO) -> str:
        from docx import Document
        try:
            doc = Document(file)
        except Exception as e:
            raise ValueError(
                f"Word 文件解析失敗（{e}）。請確認檔案未損毀，或以 Word 開啟後另存為 .docx 再上傳。"
            )
        return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())

    def _load_odt(self, file: BinaryIO) -> str:
        """ODF 文字文件 (.odt) — odfpy 抽段落/標題（含表格內 P）。公文常用 ODF。"""
        import io as _io

        from odf import teletype
        from odf.opendocument import load as _odf_load
        from odf.text import H, P
        try:
            doc = _odf_load(_io.BytesIO(file.read()))
        except Exception as e:
            raise ValueError(f"ODF 文件解析失敗（{e}）。請確認檔案未損毀或另存為 .docx 再上傳。")
        parts = [
            teletype.extractText(el)
            for el in (doc.getElementsByType(H) + doc.getElementsByType(P))
        ]
        return "\n".join(t for t in parts if t and t.strip())

    def _load_ods(self, file: BinaryIO) -> str:
        """ODF 試算表 (.ods) — odfpy 逐列逐格抽文字。"""
        import io as _io

        from odf import teletype
        from odf.opendocument import load as _odf_load
        from odf.table import Table, TableCell, TableRow
        try:
            doc = _odf_load(_io.BytesIO(file.read()))
        except Exception as e:
            raise ValueError(f"ODF 試算表解析失敗（{e}）。")
        lines = []
        for tbl in doc.getElementsByType(Table):
            for row in tbl.getElementsByType(TableRow):
                cells = [teletype.extractText(c) for c in row.getElementsByType(TableCell)]
                line = "\t".join(c for c in cells if c and c.strip())
                if line.strip():
                    lines.append(line)
        return "\n".join(lines)

    def _load_doc_legacy(self, file: BinaryIO) -> str:
        """舊版 .doc (OLE2) 用 antiword 解析。"""
        import subprocess, tempfile, os
        with tempfile.NamedTemporaryFile(suffix=".doc", delete=False) as tmp:
            tmp.write(file.read())
            tmp_path = tmp.name
        try:
            result = subprocess.run(
                ["antiword", "-w", "0", "-m", "UTF-8.txt", tmp_path],
                capture_output=True, timeout=60,
            )
            if result.returncode != 0:
                stderr = result.stderr.decode("utf-8", errors="replace").strip()
                raise ValueError(
                    f"antiword 解析失敗：{stderr or '未知錯誤'}。"
                    f"建議以 Word 開啟後另存為 .docx 再上傳。"
                )
            return result.stdout.decode("utf-8", errors="replace")
        finally:
            try: os.unlink(tmp_path)
            except OSError: pass

    def _load_excel(self, file: BinaryIO) -> str:
        import openpyxl
        wb = openpyxl.load_workbook(file, read_only=True)
        rows = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    rows.append("\t".join(cells))
        return "\n".join(rows)

    def _load_xls_legacy(self, file: BinaryIO) -> str:
        """v5.9.19: 舊版 .xls (OLE2) 用 xlrd 2.x 解析。"""
        try:
            import xlrd
        except ImportError as e:
            raise ValueError(
                "舊版 Excel (.xls) 解析需 xlrd library — 環境未安裝。"
                "請以 Excel 開啟後另存為 .xlsx 再上傳。"
            ) from e
        data = file.read()
        # BIFF8 (.xls Excel97+) 字串走 UTF-16，xlrd 自解；
        # 但 BIFF5 / codepage 缺失的舊台灣公文檔，xlrd 預設可能解不出中文 →
        # 依序 fallback encoding_override（cp950 / big5 / gb18030）。
        wb = None
        last_err: Exception | None = None
        for override in (None, "cp950", "big5", "gb18030"):
            try:
                kwargs = {"file_contents": data}
                if override is not None:
                    kwargs["encoding_override"] = override
                wb = xlrd.open_workbook(**kwargs)
                break
            except Exception as e:  # noqa: BLE001 — 換編碼重試
                last_err = e
        if wb is None:
            raise ValueError(
                f"舊版 Excel (.xls) 解析失敗：{last_err}。"
                "請確認檔案未損毀，或以 Excel 開啟後另存為 .xlsx 再上傳。"
            )
        rows = []
        for sheet in wb.sheets():
            for r_idx in range(sheet.nrows):
                cells = []
                for c in sheet.row_values(r_idx):
                    if c == "" or c is None:
                        continue
                    cells.append(str(c))
                if cells:
                    rows.append("\t".join(cells))
        return "\n".join(rows)

    def _load_csv(self, file: BinaryIO) -> str:
        import csv
        content = self._decode_bytes(file.read())
        reader = csv.reader(io.StringIO(content))
        return "\n".join("\t".join(row) for row in reader)

    def _load_text(self, file: BinaryIO) -> str:
        return self._decode_bytes(file.read())

    # ── v5.13: 音檔 / 字幕 ───────────────────────────────────────────
    @staticmethod
    def _fmt_ts(seconds: float) -> str:
        """秒 → [hh:]mm:ss 時間碼。"""
        s = int(seconds)
        h, rem = divmod(s, 3600)
        m, sec = divmod(rem, 60)
        return f"{h:d}:{m:02d}:{sec:02d}" if h else f"{m:d}:{sec:02d}"

    def _load_audio(self, file: BinaryIO) -> str:
        """音檔 → 地端 ASR 逐字稿。每段前綴 [mm:ss] 時間碼 → 檢索命中可定位音檔時刻。

        ASR 不可用（記憶體不足/載入失敗）會 raise，由 process_document 標 error 並回明確訊息。
        """
        from app.core.asr import transcribe
        segments = transcribe(file.read())
        if not segments:
            return ""
        lines = [f"[{self._fmt_ts(s['start'])}] {s['text']}" for s in segments]
        return "\n".join(lines)

    def _load_srt(self, file: BinaryIO) -> str:
        return self._parse_subtitles(self._decode_bytes(file.read()))

    def _load_vtt(self, file: BinaryIO) -> str:
        return self._parse_subtitles(self._decode_bytes(file.read()))

    @staticmethod
    def _parse_subtitles(text: str) -> str:
        """解析 SRT / WebVTT → 每段 [mm:ss] 文字。容忍兩種格式（時間分隔 , 或 .）。"""
        import re
        # 時間軸行：00:00:01,000 --> 00:00:04,000（SRT）或 00:00:01.000 --> ...（VTT）
        ts_re = re.compile(
            r"(\d{1,2}):(\d{2}):(\d{2})[.,](\d{3})\s*-->\s*\d{1,2}:\d{2}:\d{2}[.,]\d{3}"
        )
        lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
        out: list[str] = []
        i = 0
        while i < len(lines):
            m = ts_re.search(lines[i])
            if not m:
                i += 1
                continue
            h, mm, ss, _ms = (int(x) for x in m.groups())
            total = h * 3600 + mm * 60 + ss
            stamp = f"{h:d}:{mm:02d}:{ss:02d}" if h else f"{mm:d}:{ss:02d}"
            # 收集時間軸之後、到下一個空行/時間軸前的文字行
            i += 1
            buf: list[str] = []
            while i < len(lines) and lines[i].strip() and not ts_re.search(lines[i]):
                cue = lines[i].strip()
                # 跳過純數字的 SRT 序號行
                if not cue.isdigit():
                    cue = re.sub(r"<[^>]+>", "", cue)   # 去掉 VTT 內嵌標籤
                    buf.append(cue)
                i += 1
            if buf:
                out.append(f"[{stamp}] {' '.join(buf)}")
        return "\n".join(out)


class TextSplitter:
    """針對繁體中文優化的遞迴分段器。"""

    def __init__(self, chunk_size: int = 512, chunk_overlap: int = 64):
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        # 依優先順序的分段符號 (繁體中文適用)
        self.separators = ["\n\n", "\n", "。", "！", "？", "；", " ", ""]

    def split(self, text: str) -> list[str]:
        return self._split_recursive(text.strip(), self.separators)

    def _split_recursive(self, text: str, separators: list[str]) -> list[str]:
        if len(text) <= self.chunk_size:
            return [text] if text else []

        separator = separators[0] if separators else ""
        splits = text.split(separator) if separator else list(text)

        chunks: list[str] = []
        current = ""

        for split in splits:
            candidate = current + (separator if current else "") + split
            if len(candidate) <= self.chunk_size:
                current = candidate
            else:
                if current:
                    chunks.append(current)
                if len(split) > self.chunk_size and len(separators) > 1:
                    chunks.extend(self._split_recursive(split, separators[1:]))
                    current = ""
                else:
                    current = split

        if current:
            chunks.append(current)

        # 加入重疊
        if self.chunk_overlap > 0 and len(chunks) > 1:
            chunks = self._apply_overlap(chunks)

        return chunks

    def _apply_overlap(self, chunks: list[str]) -> list[str]:
        overlapped = [chunks[0]]
        for i in range(1, len(chunks)):
            prev_tail = chunks[i - 1][-self.chunk_overlap:]
            overlapped.append(prev_tail + chunks[i])
        return overlapped
